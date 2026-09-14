from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
import pathlib, os

SHOTS = pathlib.Path("/home/mdalton/Documents/JARVIS/desktop/shots")
OUT = "/tmp/NameOS-Interface-Breakdown.docx"

d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11)

def h(text, level=1):
    p = d.add_heading(text, level=level)
    return p

def para(text, bold=False, italic=False):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    return p

def bullets(items):
    for i in items:
        d.add_paragraph(i, style="List Bullet")

def shot(name, caption, width=6.3):
    f = SHOTS / ("doc-%s.png" % name)
    if not f.is_file():
        para("[screenshot missing: %s]" % name, italic=True); return
    d.add_picture(str(f), width=Inches(width))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

# ---------------------------------------------------------------- title
d.add_heading("NameOS — interface breakdown", 0)
para("For a designer. Written 27 August 2026 from the shipping build.", italic=True)
para("vDesktop, LLC", italic=True)

para("")
_b = d.add_paragraph()
_r = _b.add_run("CONFIDENTIAL — vDesktop, LLC")
_r.bold = True
_r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
para("This document and every screenshot in it are Confidential Information of "
     "vDesktop, LLC, disclosed under a signed Non-Disclosure and Work Product "
     "Agreement. Do not copy, forward, post or publish it, and do not upload it to any "
     "third-party service. If you received it without signing that agreement, please "
     "return it and tell us.", italic=True)

para("")
para("What this is for", bold=True)
para("Every screen in the product, what it is meant to do, and what we already know is "
     "wrong with it. Screenshots are the real running app at 2x, not mockups. Nothing here "
     "is aspirational — if it is in this document it is on screen today.")

para("")
para("The one-line pitch", bold=True)
para("Give it a name and it's yours. An AI crew that works on your own files, on your own "
     "computer, using your own AI account — and shows you every step it takes.")

# ---------------------------------------------------------------- constraints
h("Constraints a designer needs to know first", 1)
para("These are not preferences. They are what the thing is built out of.")
bullets([
  "Windows desktop app. It renders in WebView2 — effectively Edge — so modern CSS is fine, "
  "but there is no browser chrome and no tabs.",
  "The entire interface is ONE hand-written HTML file. No framework, no build step, no npm, "
  "no CDN. That is deliberate: the app must work with no internet and cannot go blank because "
  "someone else's server moved. Deliverables as CSS and markup, please — not a React component "
  "library.",
  "No external fonts or images. Anything used has to be inlined or system-available.",
  "The window is borderless. The top bar is the drag handle and carries its own "
  "minimise / maximise / close.",
  "Dark only, today. There is no light theme and nobody has asked for one.",
])

para("")
para("LOCKED: the ring stays.", bold=True)
para("The animated ring behind everything is not up for redesign. It was proposed for removal "
     "and the owner said no. Anything that reclaims that space has to work AROUND it, not "
     "instead of it.")

# ---------------------------------------------------------------- tokens
h("Design tokens in use today", 1)
para("The app and the marketing site share these. Changing one means changing both.")
rows = [
  ("Background", "#0B0D10", "The window and every panel sits on this"),
  ("Panel", "#12161B", "Cards, sheets, the chat panel"),
  ("Line", "#1E252D", "Every border and divider"),
  ("Text", "#DFE5EC", "Primary text"),
  ("Dim", "#7D8894", "Secondary text, labels, captions"),
  ("Accent", "#52FFA5", "Mint. Interactive things only — 'press this'"),
  ("Accent hover", "#6BFDFF", "Cyan"),
  ("On accent", "#06231A", "Text ON a mint fill"),
  ("Good", "#22C55E", "Earned green. NOT the accent, deliberately"),
  ("Warning", "#F59E0B", "Amber — a key exists but has never been tested"),
  ("Idle", "#64748B", "Grey — nothing set up"),
  ("Bad", "#FF7A7A", "Errors"),
]
t = d.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
hdr = t.rows[0].cells
hdr[0].text = "Token"; hdr[1].text = "Value"; hdr[2].text = "Where it is used"
for a,b,c in rows:
    cells = t.add_row().cells
    cells[0].text = a; cells[1].text = b; cells[2].text = c

para("")
para("Why status colours are separate from the accent", bold=True)
para("The accent means 'press this'. Green, amber and grey mean 'this is true about the "
     "world'. When a brand colour also means success, every success looks like a button and "
     "every button looks like good news. Please keep them apart.")

# ---------------------------------------------------------------- surfaces
h("The surfaces", 1)

h("1. The main window", 2)
shot("01-main", "The window as it opens: rail on the left, chat panel bottom right, ring behind everything.")
para("Structure:", bold=True)
bullets([
  "Top bar — product name, run state (Idle / Working), the connection chip on the right, "
  "then the window buttons. The whole bar is the drag handle.",
  "Left rail — icons with always-on labels, grouped by hairline rules into: what you are "
  "doing now (New chat, Folder), what it can do (Apps, Skills), how it behaves (About you, "
  "Read aloud, Permissions), then Tour, then Account pinned at the bottom.",
  "Chat panel — docked bottom right. Holds the conversation, the input, the mic, the attach "
  "button, and the working-folder chip.",
  "The ring — fills the window, reacts to speech and activity.",
])
para("Known problems:", bold=True)
bullets([
  "The chat panel is about 440px in a 1280px window. Most of the screen carries no "
  "information, and the panel crops the ring it sits on.",
  "The rail labels are one register now, but 'Read aloud' is a verb phrase among nouns.",
  "There is no visible keyboard language anywhere except Esc on sheets.",
])

h("2. The empty state — what a new person sees", 2)
para("This is the most important screen in the product and the one most worth a designer's "
     "attention. It is the first thing after install.")
bullets([
  "A sentence, four clickable starter prompts, and a link to the tour.",
  "Clicking a starter fills the box but does NOT send — the first thing somebody clicks "
  "should not be a request they have not read.",
])

h("3. Listening", 2)
shot("02-listening", "When the wake word is armed: the prompt sits top-centre and the text box hides.")
bullets([
  "'Say “Atlas” to talk to it' appears top-centre, over the ring, using whatever name "
  "the person chose.",
  "While listening, the text box is hidden and a small 'Type instead' brings it back for the "
  "rest of the session.",
  "A blocked microphone turns this bar red with an 'Open settings' button rather than being an "
  "error in the conversation.",
])
para("Open question for the designer: hiding the text box is what the owner asked for, but "
     "standard guidance calls a buried text fallback an anti-pattern. Worth a second opinion.", italic=True)

h("4. Apps — the connections panel", 2)
shot("03-apps", "Six recommended services by default; All and Connected filters; search; earned states below.")
bullets([
  "Twenty-one hosted services, six shown by default, the rest behind 'All'. Search always "
  "looks at all of them.",
  "Real brand marks, not letters. One service has no mark available and falls back to a letter.",
  "Connect is outlined at rest and fills on hover — twenty solid buttons in one grid meant "
  "nothing was the point.",
  "Below the grid: what is actually connected, including things set up outside NameOS, each "
  "with the account name and a Disconnect.",
])
para("Known problems:", bold=True)
bullets([
  "The Connected list is below the fold in a short window, and it is the most trust-building "
  "content in the app.",
  "Two ways in (rail icon and the top-bar chip) with one panel — fine, but the chip reads as "
  "a status rather than a door.",
])

h("5. Skills", 2)
shot("04-skills", "Starter skills to adapt, the ones you have written, and the two ways to add more.")
bullets([
  "Four starter skills, each described by what it does.",
  "'Write a new skill' is the primary filled button; 'Import from GitHub' is secondary.",
  "Importing shows what is in a repository and installs nothing until something is chosen.",
])

h("6. About you / About me", 2)
shot("05-aboutyou", "About you: who you are, what you are working toward, the one thing to always remember.")
shot("06-aboutme", "About me: the name, the wake word, the personality, and the voice picker.")
bullets([
  "Two tabs in one sheet. The heading follows the tab.",
  "'About me' holds the assistant's identity — this is where the product's whole promise "
  "(you name it) actually happens, and it is currently two clicks deep behind a tab.",
])
para("Biggest opportunity in the product:", bold=True)
para("Naming it is the pitch, and today it is a text field on the second tab of a settings "
     "sheet. A first-run moment built around naming would be worth more than anything else in "
     "this document.")

h("7. The first-run tour", 2)
shot("07-tour", "Seven stops. It dims the window and rings one real control at a time.")
bullets([
  "Seven stops over the real controls, click anywhere to advance, Escape or Skip to leave.",
  "Re-openable from the rail so a dismissed tour is not gone forever.",
])

# ---------------------------------------------------------------- ask
h("What we would like from a designer", 1)
bullets([
  "The first-run experience end to end: install, name it, first useful thing. Including the "
  "naming moment.",
  "What to do with the empty two thirds of the window, given the ring stays.",
  "A type scale and a spacing scale. There is not really one today — sizes were chosen "
  "individually and it shows.",
  "Iconography. The rail icons are hand-drawn SVG paths and they are inconsistent in weight.",
  "A second opinion on the listening state hiding the text box.",
])
para("")
para("What we do NOT need", bold=True)
bullets([
  "A component library or a design-system tool output. One HTML file, remember.",
  "A light theme. Nobody has asked.",
  "Anything that replaces or shrinks the ring.",
])

para("")
para("Screens in this document were captured from the shipping build on 27 August 2026 at "
     "twice the display resolution. The data in them is neutral sample data, not a real "
     "person's.", italic=True)

# Same marking in the footer, so a printed or forwarded page still carries it.
for _s in d.sections:
    _f = _s.footer.paragraphs[0]
    _f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fr = _f.add_run("CONFIDENTIAL — vDesktop, LLC — disclosed under NDA — NameOS interface breakdown, 27 Aug 2026")
    _fr.font.size = Pt(7.5)
    _fr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

d.save(OUT)
print("saved", OUT, os.path.getsize(OUT), "bytes")
